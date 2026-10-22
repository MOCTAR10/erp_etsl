"""Exports transverses typés (EX-G-08) : CSV, XLSX et PDF à partir de lignes `dict` et de colonnes `(clé, en-tête)`.

- XLSX **typé** : nombres et montants en cellules numériques (format `# ##0.00`), dates en dates, booléens « Oui/Non »
  (et non tout en texte) ; en-tête figé, filtre automatique, largeur de colonne ajustée.
- PDF : A4 paysage, en-tête répété à chaque page, métadonnées (titre, auteur, sujet), pied de page daté et paginé.
  Le profil PDF/A (ICC, XMP) n'est pas produit par ReportLab : réserve signalée (question ETSL sur la norme exigée).
- CSV : séparateur `;`, UTF-8 avec BOM (ouverture directe dans Excel).

Le téléchargement est journalisé par `users.journal.JournalExportsMiddleware` (pièce jointe).
"""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _texte(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Oui" if v else "Non"
    if isinstance(v, datetime):
        return timezone.localtime(v).strftime("%d/%m/%Y %H:%M") if timezone.is_aware(v) else v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, Decimal):
        return f"{v:.2f}".replace(".", ",") if v != v.to_integral_value() else str(v.quantize(Decimal(1)))
    if isinstance(v, (list, tuple)):
        return ", ".join(_texte(x) for x in v)
    return str(v)


def _cellule(v):
    """Valeur typée pour openpyxl (nombre, date, texte)."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "Oui" if v else "Non"
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return timezone.localtime(v).replace(tzinfo=None) if timezone.is_aware(v) else v
    if isinstance(v, (int, float, date)):
        return v
    if isinstance(v, (list, tuple, dict)):
        return _texte(v) if not isinstance(v, dict) else str(v)
    return str(v)


def to_csv(lignes, colonnes):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow([h for _, h in colonnes])
    for ligne in lignes:
        writer.writerow([_texte(ligne.get(k)) for k, _ in colonnes])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(lignes, colonnes, titre="Export"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (titre or "Export")[:31].replace("/", "-")
    ws.append([h for _, h in colonnes])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="12308F")
    for ligne in lignes:
        ws.append([_cellule(ligne.get(k)) for k, _ in colonnes])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "# ##0.00"
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = "0"
            elif isinstance(cell.value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(cell.value, date):
                cell.number_format = "DD/MM/YYYY"
    for column_cells in ws.columns:
        largeur = max(len(_texte(c.value)) for c in column_cells) + 2
        ws.column_dimensions[column_cells[0].column_letter].width = max(10, min(largeur, 60))
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    wb.properties.title = titre
    wb.properties.creator = "ETSL ERP"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(lignes, colonnes, titre="Export", sous_titre=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    genere = timezone.localtime().strftime("%d/%m/%Y %H:%M")

    def pied(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(10 * mm, 6 * mm, f"ETSL ERP — {titre} — généré le {genere}")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
        bottomMargin=12 * mm, title=titre, author="ETSL ERP", subject=sous_titre or titre, creator="ETSL ERP",
    )
    styles = getSampleStyleSheet()
    cellule = styles["BodyText"].clone("cellule", fontSize=7, leading=8.5)
    elements = [Paragraph(titre, styles["Title"])]
    if sous_titre:
        elements.append(Paragraph(sous_titre, styles["Normal"]))
    elements.append(Spacer(1, 4 * mm))
    data = [[Paragraph(f"<b>{h}</b>", cellule) for _, h in colonnes]]
    for ligne in lignes:
        data.append([Paragraph(_texte(ligne.get(k)).replace("&", "&amp;").replace("<", "&lt;"), cellule) for k, _ in colonnes])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12308f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c7cbd6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f5fa")]),
    ]))
    elements.append(table)
    doc.build(elements, onFirstPage=pied, onLaterPages=pied)
    return buffer.getvalue()


def export_response(lignes, colonnes, fmt, titre, nom_fichier, sous_titre=""):
    """Réponse HTTP en pièce jointe pour `fmt` ∈ csv | xlsx | pdf."""
    lignes = list(lignes)
    if fmt == "xlsx":
        contenu = to_xlsx(lignes, colonnes, titre)
    elif fmt == "pdf":
        contenu = to_pdf(lignes, colonnes, titre, sous_titre)
    else:
        fmt = "csv"
        contenu = to_csv(lignes, colonnes)
    response = HttpResponse(contenu, content_type=CONTENT_TYPES[fmt])
    response["Content-Disposition"] = f'attachment; filename="{nom_fichier}.{fmt}"'
    return response
