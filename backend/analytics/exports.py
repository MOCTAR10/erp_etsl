"""Couche D — exports Excel / PDF des tableaux Direction, du reporting pétrolier, des cubes et du cash-flow.

`?export=xlsx|pdf` (le paramètre `format` est réservé par DRF). Une section = un onglet Excel / un tableau PDF.
"""

import io
import re

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMATS = ("xlsx", "pdf")
TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _texte(v):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "oui" if v else "non"
    return str(v)


def sections_indicateurs(donnees, ignorer=()):
    """Aplati un dict {module: {indicateur: valeur | dict}} en sections (titre, colonnes, lignes) à 2 colonnes."""
    sections = []
    for module, valeurs in donnees.items():
        if module in ignorer or not isinstance(valeurs, dict):
            continue
        lignes = []

        def parcourir(prefixe, d):
            for k, v in d.items():
                cle = f"{prefixe}.{k}" if prefixe else k
                if isinstance(v, dict):
                    parcourir(cle, v)
                elif not isinstance(v, list):
                    lignes.append([cle, _texte(v)])

        parcourir("", valeurs)
        if lignes:
            sections.append((module, ["Indicateur", "Valeur"], lignes))
    return sections


def _xlsx(titre, sections):
    wb = Workbook()
    wb.remove(wb.active)
    for nom, colonnes, lignes in sections:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", nom)[:31] or "Feuille")
        ws.append([titre])
        ws["A1"].font = Font(bold=True)
        ws.append(list(colonnes))
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for ligne in lignes:
            ws.append([_texte(v) for v in ligne])
        for col in ws.columns:
            largeur = max(len(str(c.value or "")) for c in list(col)[1:]) + 2 if len(list(col)) > 1 else 12
            ws.column_dimensions[col[0].column_letter].width = min(max(largeur, 10), 50)
    if not wb.sheetnames:
        wb.create_sheet("Vide").append([titre])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(titre, sections):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=titre)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(titre, styles["Title"]),
        Paragraph(f"ETSL — édité le {timezone.localtime():%d/%m/%Y %H:%M}", styles["Normal"]),
        Spacer(1, 4 * mm),
    ]
    for nom, colonnes, lignes in sections:
        story.append(Paragraph(nom, styles["Heading3"]))
        table = Table([list(colonnes)] + [[_texte(v) for v in l] for l in lignes] or [["—"]], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12308f")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
        ]))
        story += [table, Spacer(1, 4 * mm)]
    doc.build(story)
    return buf.getvalue()


def reponse_export(fmt, titre, nom_fichier, sections):
    contenu = _xlsx(titre, sections) if fmt == "xlsx" else _pdf(titre, sections)
    resp = HttpResponse(contenu, content_type=TYPES[fmt])
    resp["Content-Disposition"] = f'attachment; filename="{nom_fichier}.{fmt}"'
    return resp
