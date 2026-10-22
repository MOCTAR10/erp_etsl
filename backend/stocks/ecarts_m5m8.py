"""Écarts M5 (lot m5m8) : plan d'inventaire tournant annuel, comptage, écarts valorisés et régularisation « MAD ».

- Plan annuel : à partir d'un plan tournant (périodicité, tranche), calendrier des passages de l'année, capacité de
  comptage comparée au nombre de cycles complets attendu par ETSL (paramètre provisoire `INVENTAIRE_CYCLES_AN`),
  tranche conseillée, passages réalisés et couverture (articles comptés au moins une fois / articles du dépôt).
- Écarts valorisés : écart de quantité × coût unitaire du quant (CUMP), ligne signalée au-delà de la tolérance
  (`ECART_TOLERANCE_PCT`) ; au-delà du seuil `SEUIL_VALIDATION_ECART` (valeur absolue cumulée), la régularisation
  « MAD » (mouvements d'ajustement d'inventaire) exige une validation hiérarchique avant la clôture.
- État des écarts par plan et par année, exportable en XLSX / PDF.
La signification exacte de « MAD » reste à confirmer par ETSL (libellé provisoire : mouvement d'ajustement).
"""

import io
import math
from datetime import date, timedelta
from decimal import Decimal

from .models import Inventaire, InventaireLigne, StockQuant
from .services import get_param_stock, prix_standard

CENT = Decimal("0.01")
LIBELLE_MAD = "Régularisation MAD (ajustement d'inventaire)"


def _articles_du_depot(depot):
    from referentiels.models import Article

    ids = StockQuant.objects.filter(depot=depot, quantity__gt=0).values_list("article_id", flat=True).distinct()
    return list(Article.objects.filter(pk__in=ids).order_by("code"))


def passages_de_l_annee(plan, annee):
    """Dates des passages du plan dans l'année (pas = périodicité, ancré sur le prochain passage)."""
    pas = timedelta(days=max(plan.periodicite_jours, 1))
    debut, fin = date(annee, 1, 1), date(annee, 12, 31)
    d = plan.prochaine_date
    while d > debut:
        d -= pas
    while d < debut:
        d += pas
    dates = []
    while d <= fin:
        dates.append(d)
        d += pas
    return dates


def plan_annuel(plan, annee):
    articles = _articles_du_depot(plan.depot)
    dates = passages_de_l_annee(plan, annee)
    n = len(articles)
    capacite = len(dates) * plan.articles_par_passage
    cible = get_param_stock("INVENTAIRE_CYCLES_AN")
    cycles = (Decimal(capacite) / Decimal(n)).quantize(CENT) if n else None
    conseille = math.ceil(n * float(cible) / len(dates)) if (n and dates) else None
    inventaires = Inventaire.objects.filter(plan=plan, date__year=annee).order_by("date")
    comptes = set(
        InventaireLigne.objects.filter(
            inventaire__plan=plan, inventaire__date__year=annee,
            inventaire__statut=Inventaire.Statut.CLOTURE, comptee=True,
        ).values_list("article_id", flat=True)
    )
    realises = [
        {
            "id": str(inv.pk), "code": inv.code, "date": inv.date.isoformat(), "statut": inv.statut,
            "lignes": inv.lignes.count(), "comptees": inv.lignes.filter(comptee=True).count(),
        }
        for inv in inventaires
    ]
    return {
        "plan": plan.code, "depot": plan.depot.code, "annee": annee, "articles": n,
        "periodicite_jours": plan.periodicite_jours, "articles_par_passage": plan.articles_par_passage,
        "passages_prevus": [d.isoformat() for d in dates], "capacite": capacite,
        "cycles_prevus": None if cycles is None else str(cycles), "cycles_cibles": str(cible),
        "conforme": cycles is not None and cycles >= cible, "articles_par_passage_conseille": conseille,
        "realises": realises, "articles_comptes": len(comptes),
        "couverture_pct": str((Decimal(len(comptes)) * 100 / Decimal(n)).quantize(CENT)) if n else None,
    }


def cout_unitaire_ligne(ligne):
    quant = StockQuant.objects.filter(
        depot=ligne.inventaire.depot, article=ligne.article, lot=ligne.lot
    ).first()
    if quant and quant.quantity > 0:
        return (quant.stock_value / quant.quantity).quantize(Decimal("0.0001"))
    return Decimal(str(prix_standard(ligne.article) or 0))


def lignes_valorisees(inventaire):
    tolerance = get_param_stock("ECART_TOLERANCE_PCT")
    lignes = []
    for ligne in inventaire.lignes.select_related("article", "lot", "inventaire__depot").order_by("article__code"):
        ecart = ligne.ecart
        pct = (abs(ecart) * 100 / ligne.quantite_systeme).quantize(CENT) if ligne.quantite_systeme else None
        cu = cout_unitaire_ligne(ligne)
        lignes.append({
            "id": str(ligne.pk), "article": ligne.article.code, "article_label": ligne.article.label,
            "lot": ligne.lot.code if ligne.lot else None, "comptee": ligne.comptee,
            "quantite_systeme": ligne.quantite_systeme, "quantite_reelle": ligne.quantite_reelle,
            "ecart": ecart, "ecart_pct": pct,
            "hors_tolerance": bool(ecart) and (pct is None or pct > tolerance),
            "cout_unitaire": cu, "valeur_ecart": (ecart * cu).quantize(CENT),
        })
    return lignes


def synthese_ecarts(inventaire):
    lignes = lignes_valorisees(inventaire)
    total_abs = sum((abs(l["valeur_ecart"]) for l in lignes), Decimal("0"))
    seuil = get_param_stock("SEUIL_VALIDATION_ECART")
    return {
        "lignes": lignes,
        "valeur_nette": sum((l["valeur_ecart"] for l in lignes), Decimal("0")),
        "valeur_absolue": total_abs, "seuil_validation": seuil,
        "validation_requise": total_abs > seuil,
        "validee": inventaire.ecarts_valides_le is not None,
    }


def controler_regularisation(inventaire):
    """Refuse la clôture si les écarts valorisés dépassent le seuil sans validation de la régularisation MAD."""
    if inventaire.ecarts_valides_le is not None:
        return
    s = synthese_ecarts(inventaire)
    if s["validation_requise"]:
        raise ValueError(
            f"Écarts valorisés {s['valeur_absolue']:.0f} FCFA au-delà du seuil {s['seuil_validation']:.0f} FCFA : "
            "la régularisation MAD doit être validée avant la clôture."
        )


def etat_ecarts(plan, annee):
    rows = []
    for inv in Inventaire.objects.filter(plan=plan, date__year=annee).order_by("date", "code"):
        for l in lignes_valorisees(inv):
            if not l["ecart"]:
                continue
            rows.append({"inventaire": inv.code, "date": inv.date.isoformat(), "statut": inv.statut, **l})
    return {
        "plan": plan.code, "depot": plan.depot.code, "annee": annee, "rows": rows,
        "valeur_nette": sum((r["valeur_ecart"] for r in rows), Decimal("0")),
        "valeur_absolue": sum((abs(r["valeur_ecart"]) for r in rows), Decimal("0")),
        "hors_tolerance": sum(1 for r in rows if r["hors_tolerance"]),
    }


def en_texte(data):
    """Décimaux en chaînes pour une `Response` brute."""
    if isinstance(data, dict):
        return {k: en_texte(v) for k, v in data.items()}
    if isinstance(data, list):
        return [en_texte(v) for v in data]
    if isinstance(data, Decimal):
        return str(data)
    return data


# ---------------------------------------------------------------------------------------------------------------
# Exports tabulaires (XLSX / PDF) partagés par les états M5 et M8
# ---------------------------------------------------------------------------------------------------------------
def tableau_xlsx(titre, entetes, lignes):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws.append([titre])
    ws["A1"].font = Font(bold=True)
    ws.append(entetes)
    for c in ws[2]:
        c.font = Font(bold=True)
    for ligne in lignes:
        ws.append(["" if v is None else (float(v) if isinstance(v, Decimal) else v) for v in ligne])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tableau_pdf(titre, entetes, lignes, sous_titre=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=titre)
    styles = getSampleStyleSheet()
    elements = [Paragraph(titre, styles["Title"])]
    if sous_titre:
        elements.append(Paragraph(sous_titre, styles["Normal"]))
    elements.append(Spacer(1, 4 * mm))
    data = [entetes] + [["" if v is None else str(v) for v in ligne] for ligne in lignes]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12308f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
    ]))
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()


def reponse_export(fmt, nom, titre, entetes, lignes, sous_titre=""):
    from django.http import HttpResponse

    if fmt == "xlsx":
        contenu = tableau_xlsx(titre, entetes, lignes)
        ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        contenu = tableau_pdf(titre, entetes, lignes, sous_titre)
        ctype = "application/pdf"
    resp = HttpResponse(contenu, content_type=ctype)
    resp["Content-Disposition"] = f'attachment; filename="{nom}.{fmt}"'
    return resp
