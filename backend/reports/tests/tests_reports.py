"""Tests Increment 10 : exports (RF-98/74), SAGE (RF-68/80), dashboard (RF-94/95/96), rétention (RF-50/51/52)."""

import io
from datetime import date, timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from documents.models import Document, DocumentType
from workflow.models import Circuit, CircuitStep, Task

User = get_user_model()


@override_settings(
    STORAGES={"default": {"BACKEND": "django.core.files.storage.InMemoryStorage"}}
)
class ReportsTests(APITestCase):
    def setUp(self):
        self.comptable = User.objects.create_user(
            email="compta@etls.local", password="MotDePasse#2026", role=User.Role.COMPTABLE
        )
        self.direction = User.objects.create_user(
            email="direction@etls.local", password="MotDePasse#2026", role=User.Role.DIRECTION
        )
        self.admin = User.objects.create_user(
            email="admin@etls.local", password="MotDePasse#2026", role=User.Role.ADMIN, is_staff=True
        )
        self.facture_type = DocumentType.objects.create(
            code="factures_fournisseurs", label="Factures fournisseurs", retention_years=10
        )
        self.devis_type = DocumentType.objects.create(
            code="devis", label="Devis", retention_years=5
        )
        self.circuit = Circuit.objects.create(code="c1", label="Compta", max_days=5)

    def _auth(self, user):
        tokens = self.client.post(
            "/api/users/token/",
            {"email": user.email, "password": "MotDePasse#2026"},
            format="json",
        ).data
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {tokens['access']}")

    def _doc(self, title="Facture auditable", doc_type=None, doc_date=None, amount=None):
        self._auth(self.comptable)
        payload = {
            "title": title,
            "type": (doc_type or self.facture_type).id,
            "document_date": doc_date or "2024-05-12",
            "file": SimpleUploadedFile("f.pdf", b"contenu", content_type="application/pdf"),
        }
        if amount is not None:
            payload["amount"] = amount
        resp = self.client.post(
            "/api/documents/documents/", payload, format="multipart"
        )
        return resp.data["id"]

    def test_export_csv(self):
        self._doc()
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/export/?export=csv")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp["Content-Type"], "text/csv")
        body = resp.content.decode("utf-8")
        self.assertIn("titre", body)
        self.assertIn("Facture auditable", body)

    def test_export_xlsx(self):
        self._doc()
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/export/?export=xlsx")
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content[:400])
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws["A1"].value, "id")
        self.assertIn("Facture auditable", [c.value for c in ws[2]])

    def test_export_pdf(self):
        self._doc()
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/export/?export=pdf")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_export_respects_visibility(self):
        from documents.models import Dossier, DossierAccess

        dossier = Dossier.objects.create(name="Dossier secret", created_by=self.admin)
        DossierAccess.objects.create(
            dossier=dossier, user=self.comptable,
            permission=DossierAccess.Permission.DENY, granted_by=self.admin,
        )
        self._auth(self.admin)
        resp = self.client.post(
            "/api/documents/documents/",
            {
                "title": "Confidentielle",
                "type": self.facture_type.id,
                "dossier": dossier.id,
                "file": SimpleUploadedFile("f.pdf", b"c", content_type="application/pdf"),
            },
            format="multipart",
        )
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/export/?export=csv")
        body = resp.content.decode("utf-8")
        self.assertNotIn("Confidentielle", body)
        self._auth(self.admin)
        resp = self.client.get("/api/reports/export/?export=csv")
        self.assertIn("Confidentielle", resp.content.decode("utf-8"))

    def test_sage_export(self):
        self._auth(self.comptable)
        doc_id = self.client.post(
            "/api/documents/documents/",
            {
                "title": "Facture EDF",
                "type": self.facture_type.id,
                "counterparty": "EDF SA",
                "document_date": "2024-05-12",
                "amount": "125000",
                "file": SimpleUploadedFile("f.pdf", b"contenu", content_type="application/pdf"),
            },
            format="multipart",
        ).data["id"]
        # SAGE ne concerne que les écritures validées/archivées (RF-68/80).
        Document.objects.filter(id=doc_id).update(status=Document.Status.ARCHIVED)
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/sage/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        body = resp.content.decode("utf-8")
        self.assertIn("EDF SA", body)
        self.assertIn("125000", body)

    def test_dashboard_kpis(self):
        self._doc("Facture EDF", amount="100")
        self._doc("Devis", doc_type=self.devis_type, amount="50")
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/dashboard/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["documents"]["total"], 2)
        self.assertIn("draft", resp.data["documents"]["by_status"])

    def test_retention_due_and_bordereau(self):
        old = self._doc("Vieille facture", doc_date="2010-01-01")
        recent = self._doc("Facture récente", doc_date="2025-01-01")
        self._auth(self.comptable)
        resp = self.client.get("/api/reports/retention/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        due = [str(d["id"]) for d in resp.data["due"]]
        self.assertIn(old, due)
        self.assertNotIn(recent, due)

        resp = self.client.get("/api/reports/retention/?export=csv")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp["Content-Type"], "text/csv")
        body = resp.content.decode("utf-8")
        self.assertIn("bordereau", resp["Content-Disposition"].lower())
        self.assertIn("Vieille facture", body)

    def test_retention_check_command(self):
        from django.core.management import call_command

        self._auth(self.comptable)
        self.client.post(
            "/api/documents/documents/",
            {
                "title": "Facture 2010",
                "type": self.facture_type.id,
                "document_date": "2010-01-01",
                "file": SimpleUploadedFile("f.pdf", b"c", content_type="application/pdf"),
            },
            format="multipart",
        )
        out = io.StringIO()
        call_command("retention_check", stdout=out)
        self.assertIn("1 document(s)", out.getvalue())

    def test_dashboard_admin_and_role_access(self):
        self._doc()
        self._auth(self.comptable)
        self.assertEqual(
            self.client.get("/api/reports/dashboard/").status_code, status.HTTP_200_OK
        )
