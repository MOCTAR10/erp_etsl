"""Écarts M3 (lot 2) : paramètres provisoires, croisement pointages ↔ M9, besoins matière, avancement, marge à
terminaison, indicateurs, plan de charge journalier et machine, escalades, impressions PDF, rapports périodiques."""

import io
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q, Sum
from django.utils import timezone

from .models import (
    BesoinMatiere,
    CapaciteAtelier,
    EscaladeSurcharge,
    OrdreFabrication,
    ParametreOperations,
    PointageChantier,
    RapportExecution,
)
from config.i18n import en_francais, tr

PARAMETRES_OPERATIONS_DEFAUT = {
    "PLANNING_HORIZON_SEMAINES": (
        "Planning retenu : Kanban des OF + plan de charge hebdomadaire et journalier (sans Gantt ni ordonnancement "
        "automatique) — horizon affiché", "8", "semaines"),
    "JOURS_OUVRES_SEMAINE": ("Jours ouvrés par semaine (vue journalière du plan de charge)", "5", "jours"),
    "SEUIL_ESCALADE_GR_PCT": (
        "Surcharge au-delà de laquelle la direction est saisie (sous-traitance / location GLOBAL RENTAL)", "20", "%"),
    "PLAFOND_HEURES_JOUR": ("Plafond d'heures pointées par salarié et par jour, tous OF confondus", "12", "h"),
    "METHODE_AVANCEMENT": (
        "Calcul du % d'avancement : 1 = heures pointées / prévues, 2 = quantités réalisées / commandées, "
        "3 = dernier % déclaré (rapport d'avancement)", "1", "code"),
    "TAUX_HORAIRE_DEFAUT": (
        "Taux de valorisation des heures non encore valorisées par la paie (marge à terminaison)", "6000", "FCFA/h"),
    "RETENTION_OPERATIONS_ANS": ("Conservation des données opérationnelles (OF, pointages, dossiers, rapports)", "10", "ans"),
    "MULTI_AFFAIRES_ACTIF": (
        "Multi-affaires : 1 = un OF commun peut être réparti entre plusieurs affaires (clé en %), 0 = un OF = une affaire",
        "1", "0 / 1"),
    "RECU_CLIENT_OBLIGATOIRE": (
        "Facturation d'une situation : 1 = reçu client obligatoire avant facturation, 0 = facultatif", "0", "0 / 1"),
    # Lot a3 (hors ligne, CL-OM-08) : périmètre provisoire = toutes les saisies terrain du CDC.
    "TERRAIN_DELAI_SYNCHRO_JOURS": (
        "Hors ligne : une saisie synchronisée plus de N jours après sa saisie est mise en conflit (0 = sans limite)",
        "7", "jours"),
    "TERRAIN_DOUBLON_CONFLIT": (
        "Hors ligne : 1 = un pointage identique (OF, opérateur, jour, heures) déjà au serveur est mis en conflit, "
        "0 = accepté", "1", "0 / 1"),
}
PARAMETRES_BOOLEENS = ("MULTI_AFFAIRES_ACTIF", "RECU_CLIENT_OBLIGATOIRE", "TERRAIN_DOUBLON_CONFLIT")
ACTIFS = (OrdreFabrication.Status.PREVU, OrdreFabrication.Status.LANCE, OrdreFabrication.Status.EN_COURS)
CLOS = (OrdreFabrication.Status.CLOTURE, OrdreFabrication.Status.ANNULE)
ZERO = Decimal("0")


def _q(v, pas="0.01"):
    return None if v is None else Decimal(v).quantize(Decimal(pas))


def _s(v):
    return None if v is None else str(_q(v))


def assurer_parametres():
    for code, (label, valeur, unite) in PARAMETRES_OPERATIONS_DEFAUT.items():
        ParametreOperations.objects.get_or_create(code=code, defaults={"label": label, "valeur": Decimal(valeur), "unite": unite})


def param(code):
    label, valeur, unite = PARAMETRES_OPERATIONS_DEFAUT[code]
    obj, _ = ParametreOperations.objects.get_or_create(code=code, defaults={"label": label, "valeur": Decimal(valeur), "unite": unite})
    return obj.valeur


def parametres_provisoires():
    assurer_parametres()
    return ParametreOperations.objects.filter(statut_valeur=ParametreOperations.StatutValeur.PROVISOIRE).exists()


def _destinataires(*roles, ordre=None):
    from django.contrib.auth import get_user_model

    if ordre is not None and ordre.responsible_id:
        return [ordre.responsible]
    return list(get_user_model().objects.filter(is_active=True, role__in=roles))


def _notifier(destinataires, sujet, message, kind="escalation"):
    from workflow.notifications import notify

    for user in destinataires:
        notify(user, sujet, message, kind=kind)


# ---------------------------------------------------------------------------------------------------------------
# Croisement avec M9 : absences / congés, plafond journalier, affectations
# ---------------------------------------------------------------------------------------------------------------
def employe_du_pointage(pointage):
    from rh_paie.models import Employe

    if pointage.employe_id:
        return pointage.employe
    return Employe.objects.filter(user_id=pointage.worker_id).first() if pointage.worker_id else None


def indisponibilite(employe, debut, fin=None):
    """Motif d'indisponibilité M9 (absence validée, congé approuvé) chevauchant [debut, fin], sinon None."""
    from rh_paie.models import Absence, DemandesConge

    fin = fin or debut
    absence = Absence.objects.filter(
        employe=employe, statut=Absence.Statut.VALIDEE, date_debut__lte=fin, date_fin__gte=debut
    ).first()
    if absence:
        return f"{absence.get_type_display()} du {absence.date_debut:%d/%m/%Y} au {absence.date_fin:%d/%m/%Y}"
    conge = DemandesConge.objects.filter(
        employe=employe, statut__in=[DemandesConge.Statut.APPROUVE, DemandesConge.Statut.VALIDE_RH],
        date_debut__lte=fin, date_fin__gte=debut,
    ).first()
    if conge:
        return f"{conge.get_type_display()} du {conge.date_debut:%d/%m/%Y} au {conge.date_fin:%d/%m/%Y}"
    return None


def controler_pointage_croise(employe, jour, heures, exclure=None):
    """Pointage croisé avec M9 : salarié absent / en congé ⇒ refus ; cumul du jour > plafond paramétré ⇒ refus."""
    nom = f"{employe.nom} {employe.prenom}".strip()
    motif = indisponibilite(employe, jour)
    if motif:
        raise ValidationError(tr("%(nom)s est indisponible le %(jour)s (%(motif)s) : pointage refusé (croisement M9).") % {"nom": nom, "jour": format(jour, '%d/%m/%Y'), "motif": motif})
    qs = PointageChantier.objects.filter(date=jour).filter(
        Q(employe=employe) | (Q(worker_id=employe.user_id) if employe.user_id else Q(pk__in=[]))
    )
    if exclure is not None:
        qs = qs.exclude(pk=exclure)
    cumul = qs.aggregate(t=Sum("hours"))["t"] or ZERO
    plafond = param("PLAFOND_HEURES_JOUR")
    if cumul + heures > plafond:
        raise ValidationError(
            tr("%(nom)s : %(q)s h pointées le %(jour)s pour un plafond de %(q2)s h/jour (paramètre provisoire).") % {"nom": nom, "q": _q(cumul + heures), "jour": format(jour, '%d/%m/%Y'), "q2": _q(plafond)}
        )


def absents_affectation(affectation):
    """Membres de l'équipe indisponibles (M9) sur la période de l'affectation."""
    return [
        {"employe": str(m.pk), "nom": f"{m.nom} {m.prenom}".strip(), "motif": motif}
        for m in affectation.equipe.membres.all()
        if (motif := indisponibilite(m, affectation.date_debut, affectation.date_fin))
    ]


def controler_mobilisation(affectation):
    membres = affectation.equipe.membres.count()
    if not membres:
        return None
    absents = absents_affectation(affectation)
    disponibles = membres - len(absents)
    if disponibles < affectation.effectif:
        noms = ", ".join(a["nom"] for a in absents)
        return (f"Effectif disponible insuffisant : {disponibles} / {affectation.effectif} (absents M9 sur la période : "
                f"{noms}) — mobilisation refusée.")
    return None


# ---------------------------------------------------------------------------------------------------------------
# Besoins matière : disponibilité, réservation, rupture ⇒ alerte + DA M2
# ---------------------------------------------------------------------------------------------------------------
def stock_physique(article_id):
    from stocks.models import StockQuant

    return StockQuant.objects.filter(article_id=article_id).aggregate(t=Sum("quantity"))["t"] or ZERO


def _demande_achat(besoin, manque, user):
    from achats.models import PurchaseRequest, PurchaseRequestLine

    of = besoin.ordre
    da = PurchaseRequest.objects.create(
        title=f"Réapprovisionnement {of.code} — {besoin.article.label}"[:200],
        requester=user if getattr(user, "is_authenticated", False) else None,
        department="Opérations (M3)", expected_date=of.planned_start, is_urgent=True,
        analytic=of.affaire.analytic if of.affaire_id and of.affaire.analytic_id else None,
        notes=f"DA générée par M3 : rupture matière sur l'OF {of.code} (besoin {besoin.quantite}, manque {manque}).",
    )
    PurchaseRequestLine.objects.create(request=da, article=besoin.article, label=besoin.article.label[:200], quantity=manque)
    return da


@transaction.atomic
def verifier_besoins(of, user=None):
    """Contrôle de disponibilité : stock physique − réservations des autres OF actifs. Rupture ⇒ DA + alerte."""
    resultat = []
    ruptures = []
    for besoin in BesoinMatiere.objects.select_for_update(of=("self",)).filter(ordre=of).exclude(
        statut=BesoinMatiere.Statut.SERVI
    ).select_related("article"):
        reserve_autres = BesoinMatiere.objects.filter(
            article_id=besoin.article_id, statut=BesoinMatiere.Statut.RESERVE
        ).exclude(pk=besoin.pk).exclude(ordre__status__in=CLOS).aggregate(t=Sum("quantite"))["t"] or ZERO
        dispo = stock_physique(besoin.article_id) - reserve_autres
        besoin.disponible = dispo
        besoin.verifie_le = timezone.now()
        if dispo >= besoin.quantite:
            besoin.statut = BesoinMatiere.Statut.RESERVE
        else:
            besoin.statut = BesoinMatiere.Statut.RUPTURE
            manque = besoin.quantite - max(dispo, ZERO)
            if besoin.demande_achat_id is None:
                besoin.demande_achat = _demande_achat(besoin, manque, user)
            ruptures.append(f"{besoin.article.code} (manque {_q(manque)})")
        besoin.save(update_fields=["statut", "disponible", "verifie_le", "demande_achat"])
        from stocks.reservations import liberer_besoin, reserver_besoin

        if besoin.statut == BesoinMatiere.Statut.RESERVE:
            reserver_besoin(besoin, user)  # réservation PHYSIQUE en stock (M5)
        else:
            liberer_besoin(besoin.pk, "Rupture constatée au contrôle de disponibilité")
        resultat.append(besoin)
    if ruptures:
        _notifier(
            _destinataires("chef_atelier", "directeur_operations", "logistique", ordre=of),
            f"Rupture matière — {of.code}",
            f"OF {of.code} : rupture sur {', '.join(ruptures)}. DA de réapprovisionnement générée(s) ; "
            "ré-ordonnancement de l'OF à envisager.",
        )
    return resultat


# ---------------------------------------------------------------------------------------------------------------
# Avancement, marge à terminaison, indicateurs
# ---------------------------------------------------------------------------------------------------------------
def avancement_of(of):
    """% d'avancement selon la méthode paramétrée (provisoire) ; 100 % pour un OF terminé ou clôturé."""
    if of.status in (OrdreFabrication.Status.TERMINE, OrdreFabrication.Status.CLOTURE):
        return Decimal("100")
    methode = int(param("METHODE_AVANCEMENT"))
    valides = of.rapports.filter(statut=RapportExecution.Statut.VALIDE)
    if methode == 2:
        if not of.quantity:
            return None
        fait = valides.aggregate(t=Sum("quantite_realisee"))["t"] or ZERO
        pct = fait * 100 / of.quantity
    elif methode == 3:
        dernier = valides.exclude(avancement__isnull=True).order_by("-date", "-created_at").first()
        if dernier is None:
            return ZERO
        pct = Decimal(dernier.avancement)
    else:
        if not of.calculated_hours:
            return None
        pointe = of.pointages.aggregate(t=Sum("hours"))["t"] or ZERO
        pct = pointe * 100 / of.calculated_hours
    return _q(min(pct, Decimal("100")), "0.1")


def _parts(of, multi):
    """[(affaire, part)] : répartition multi-affaires si activée et saisie, sinon l'affaire de l'OF à 100 %."""
    if multi:
        reps = list(of.repartitions.select_related("affaire"))
        if reps:
            return [(r.affaire, r.part_pct / 100) for r in reps]
    return [(of.affaire, Decimal("1"))] if of.affaire_id else []


def couts_ordres(ordres):
    """Coûts engagés par OF : main-d'œuvre chargée (paie M9, PR #49) + heures non valorisées × taux provisoire + matière."""
    from rh_paie.compta_m9b import cout_main_oeuvre
    from stocks.models import MouvementStock

    mo = {l["ordre"]: l for l in cout_main_oeuvre(date(2000, 1, 1), timezone.localdate() + timedelta(days=366))}
    taux = param("TAUX_HORAIRE_DEFAUT")
    couts = {}
    for of in ordres:
        ligne = mo.get(str(of.pk))
        pointe = of.pointages.aggregate(t=Sum("hours"))["t"] or ZERO
        transferees = Decimal(ligne["heures"]) if ligne else ZERO
        non_val = Decimal(ligne["heures_non_valorisees"]) if ligne else ZERO
        cout_mo = Decimal(ligne["cout"]) if ligne else ZERO
        heures_a_valoriser = (pointe - transferees) + non_val
        matiere = ZERO
        for m in MouvementStock.objects.filter(ordre=of, type_mouvement=MouvementStock.Type.CONSOMMATION, executed=True):
            pu = m.cout_unitaire if m.cout_unitaire is not None else (m.prix_unitaire or ZERO)
            matiere += m.quantite * pu
        couts[of.pk] = {
            "mo_paie": cout_mo, "mo_estimee": heures_a_valoriser * taux, "matiere": matiere,
            "total": cout_mo + heures_a_valoriser * taux + matiere,
        }
    return couts


def avancement_et_marges(voir_montants=True):
    """Par affaire : % chantier (pondéré par les heures prévues), coûts engagés, coût et marge à terminaison."""
    from commercial.models import Affaire

    multi = bool(param("MULTI_AFFAIRES_ACTIF"))
    ordres = list(OrdreFabrication.objects.exclude(status=OrdreFabrication.Status.ANNULE).prefetch_related("repartitions"))
    couts = couts_ordres(ordres)
    par_affaire = defaultdict(lambda: {"poids": ZERO, "pondere": ZERO, "engage": ZERO, "ordres": []})
    lignes_of = []
    for of in ordres:
        pct = avancement_of(of)
        lignes_of.append({
            "ordre": str(of.pk), "code": of.code, "label": of.label, "statut": of.status, "scope": of.scope,
            "affaire": of.affaire.code if of.affaire_id else None, "avancement_pct": _s(pct),
            "heures_prevues": _s(of.calculated_hours), "heures_pointees": _s(of.pointages.aggregate(t=Sum("hours"))["t"] or ZERO),
        })
        for affaire, part in _parts(of, multi):
            agg = par_affaire[affaire.pk]
            poids = (of.calculated_hours or Decimal("1")) * part
            agg["poids"] += poids
            agg["pondere"] += (pct or ZERO) * poids
            agg["engage"] += couts[of.pk]["total"] * part
            agg["ordres"].append(of.code)
    affaires = Affaire.objects.in_bulk(list(par_affaire))
    sortie = []
    for pk, agg in par_affaire.items():
        affaire = affaires[pk]
        pct = agg["pondere"] / agg["poids"] if agg["poids"] else ZERO
        cat = agg["engage"] * 100 / pct if pct else None
        ca = affaire.contract_amount or ZERO
        marge = ca - cat if cat is not None else None
        ligne = {
            "affaire": str(pk), "affaire_code": affaire.code, "affaire_label": affaire.title,
            "ordres": agg["ordres"], "avancement_pct": _s(_q(pct, "0.1")),
            "chiffre_affaires": _s(ca), "couts_engages": _s(agg["engage"]), "cout_a_terminaison": _s(cat),
            "marge_a_terminaison": _s(marge),
            "marge_pct": _s(marge * 100 / ca) if marge is not None and ca else None,
        }
        if not voir_montants:
            for champ in ("chiffre_affaires", "couts_engages", "cout_a_terminaison", "marge_a_terminaison", "marge_pct"):
                ligne[champ] = None
        sortie.append(ligne)
    sortie.sort(key=lambda l: l["affaire_code"])
    return {"ordres": lignes_of, "affaires": sortie, "methode": int(param("METHODE_AVANCEMENT"))}


def publier_marges(user=None):
    """Transmission vers M11 : un événement outbox `operations.marge_terminaison` par affaire (montants complets)."""
    from outbox.models import publish

    donnees = avancement_et_marges(voir_montants=True)["affaires"]
    for ligne in donnees:
        publish("operations.marge_terminaison", {**ligne, "publie_le": timezone.now().isoformat(),
                                                 "publie_par": getattr(user, "email", None)})
    return len(donnees)


def indicateurs(debut, fin):
    """Productivité (heures standard / heures pointées) et taux de rebut (rebut / produit) sur la période."""
    rapports = RapportExecution.objects.filter(statut=RapportExecution.Statut.VALIDE, date__gte=debut, date__lte=fin).select_related("ordre__gamme")
    realise = rebut = standard = ZERO
    ordres = set()
    for r in rapports:
        realise += r.quantite_realisee
        rebut += r.quantite_rebut
        ordres.add(r.ordre_id)
        if r.ordre.gamme_id:
            standard += r.ordre.gamme.total_hours_decimal * r.quantite_realisee
    pointages = PointageChantier.objects.filter(date__gte=debut, date__lte=fin)
    heures_of = pointages.filter(ordre_id__in=ordres).aggregate(t=Sum("hours"))["t"] or ZERO
    directes = pointages.filter(nature_mo=PointageChantier.NatureMo.DIRECTE).aggregate(t=Sum("hours"))["t"] or ZERO
    indirectes = pointages.filter(nature_mo=PointageChantier.NatureMo.INDIRECTE).aggregate(t=Sum("hours"))["t"] or ZERO
    jour = timezone.localdate()
    return {
        "debut": debut.isoformat(), "fin": fin.isoformat(),
        "quantite_realisee": _s(realise), "quantite_rebut": _s(rebut),
        "taux_rebut_pct": _s(rebut * 100 / (realise + rebut)) if realise + rebut else None,
        "heures_standard": _s(standard), "heures_pointees_of": _s(heures_of),
        "productivite_pct": _s(standard * 100 / heures_of) if heures_of and standard else None,
        "heures_directes": _s(directes), "heures_indirectes": _s(indirectes),
        "taux_mo_indirecte_pct": _s(indirectes * 100 / (directes + indirectes)) if directes + indirectes else None,
        "of_en_retard": OrdreFabrication.objects.filter(status__in=ACTIFS, planned_end__lt=jour).count(),
        "ruptures_ouvertes": BesoinMatiere.objects.filter(statut=BesoinMatiere.Statut.RUPTURE).exclude(ordre__status__in=CLOS).count(),
    }


# ---------------------------------------------------------------------------------------------------------------
# Plan de charge : machine, vue journalière, escalade GR
# ---------------------------------------------------------------------------------------------------------------
def heures_machine(of):
    if not of.gamme_id:
        return ZERO
    unitaire = of.gamme.operations.aggregate(t=Sum("machine_hours"))["t"] or ZERO
    return unitaire * (of.quantity or 1)


def _jours_ouvres(debut, fin, par_semaine):
    jours = []
    j = debut
    while j <= fin:
        if j.weekday() < par_semaine:
            jours.append(j)
        j += timedelta(days=1)
    return jours or [debut]


def charge_journaliere(debut=None, nb_jours=14):
    """Charge planifiée par jour ouvré (main-d'œuvre et machine) vs capacité hebdomadaire ÷ jours ouvrés."""
    debut = debut or timezone.localdate()
    fin = debut + timedelta(days=nb_jours - 1)
    par_semaine = int(param("JOURS_OUVRES_SEMAINE")) or 5
    charge = defaultdict(Decimal)
    for of in OrdreFabrication.objects.filter(status__in=ACTIFS).select_related("gamme"):
        d0 = of.planned_start or timezone.localdate()
        d1 = of.planned_end or d0
        jours = _jours_ouvres(d0, d1, par_semaine)
        for ressource, heures in (("rh", of.calculated_hours or ZERO), ("machine", heures_machine(of))):
            part = heures / len(jours)
            for j in jours:
                if debut <= j <= fin:
                    charge[(of.scope, ressource, j)] += part
    capacites = {(c.scope, c.resource): c.hours_per_week / par_semaine for c in CapaciteAtelier.objects.all()}
    lignes = []
    j = debut
    while j <= fin:
        for scope in OrdreFabrication.Scope.values:
            for ressource in CapaciteAtelier.Resource.values:
                heures = charge.get((scope, ressource, j), ZERO)
                cap = capacites.get((scope, ressource))
                if not heures and cap is None:
                    continue
                lignes.append({
                    "date": j.isoformat(), "scope": scope, "resource": ressource, "planned_hours": _s(heures),
                    "capacity": _s(cap) if cap is not None else None,
                    "overload": bool(cap is not None and j.weekday() < par_semaine and heures > cap),
                })
        j += timedelta(days=1)
    return {"debut": debut.isoformat(), "fin": fin.isoformat(), "jours_ouvres": par_semaine, "lignes": lignes}


def escalader_surcharges(user=None):
    """Surcharges au-delà du seuil provisoire : escalade unique par (périmètre, ressource, semaine) vers la direction."""
    from .cloture import alertes_surcharge

    seuil = param("SEUIL_ESCALADE_GR_PCT")
    creees = []
    for a in alertes_surcharge():
        if Decimal(str(a["overload_pct"])) <= seuil:
            continue
        esc, cree = EscaladeSurcharge.objects.get_or_create(
            scope=a["scope"], resource=a.get("resource", "rh"), semaine=date.fromisoformat(a["week"]),
            defaults={"surcharge_pct": Decimal(str(a["overload_pct"])),
                      "escalade_par": user if getattr(user, "is_authenticated", False) else None},
        )
        if cree:
            creees.append(esc)
            _notifier(
                _destinataires("directeur_operations", "dga", "pdg"),
                f"Escalade surcharge {a['scope']} — semaine du {esc.semaine:%d/%m/%Y}",
                f"Surcharge de {a['overload_pct']} % ({a['planned_hours']} h pour {a['capacity']} h, ressource "
                f"{a.get('resource', 'rh')}). Arbitrage : replanification, sous-traitance ou location GLOBAL RENTAL.",
            )
    return creees


# ---------------------------------------------------------------------------------------------------------------
# Impressions PDF et rapports
# ---------------------------------------------------------------------------------------------------------------
def _d(v):
    return f"{v:%d/%m/%Y}" if v else "—"


def _pdf(titre, blocs, tableau=None):
    from hse.services import _pdf as pdf_hse

    return pdf_hse(titre, blocs, tableau)


def of_pdf(of):
    blocs = [
        ("OF", f"{of.code} — {of.label}"), (tr("Affaire"), of.affaire.code if of.affaire_id else "—"),
        (tr("Périmètre"), of.get_scope_display()), (tr("Statut"), of.get_status_display()),
        (tr("Quantité"), of.quantity), (tr("Période prévue"), f"{_d(of.planned_start)} → {_d(of.planned_end)}"),
        (tr("Heures prévues"), of.calculated_hours), (tr("Gamme"), f"{of.gamme.code} — {of.gamme.label}" if of.gamme_id else "—"),
        (tr("Avancement"), f"{avancement_of(of) or 0} %"),
    ]
    tableau = [[tr("Rubrique"), tr("Référence"), tr("Détail"), tr("Quantité / heures")]]
    if of.gamme_id:
        for op in of.gamme.operations.all():
            tableau.append([tr("Opération"), str(op.sequence), f"{op.label} ({op.get_poste_display()})", f"{op.planned_hours} h"])
    for b in of.besoins_matiere.select_related("article"):
        tableau.append([tr("Matière"), b.article.code, f"{b.article.label[:60]} — {b.get_statut_display()}", str(b.quantite)])
    for p in of.pointages.all()[:50]:
        tableau.append([tr("Pointage"), p.code, f"{_d(p.date)} · {p.get_nature_mo_display()}", f"{p.hours} h"])
    return _pdf(tr("Ordre de fabrication %(of_code)s") % {"of_code": of.code}, blocs, tableau if len(tableau) > 1 else None)


def gamme_pdf(gamme):
    blocs = [(tr("Gamme"), f"{gamme.code} — {gamme.label}"), (tr("Statut"), gamme.get_status_display()),
             (tr("Temps total"), f"{gamme.total_hours_decimal} h")]
    tableau = [[tr("Ordre"), tr("Opération"), tr("Poste"), tr("Heures MO"), tr("Heures machine")]] + [
        [str(op.sequence), op.label, op.get_poste_display(), str(op.planned_hours), str(op.machine_hours)]
        for op in gamme.operations.all()
    ]
    return _pdf(tr("Gamme opératoire %(gamme_code)s") % {"gamme_code": gamme.code}, blocs, tableau if len(tableau) > 1 else None)


def fiche_preparation_pdf(fiche):
    of = fiche.ordre
    blocs = [(tr("Fiche"), fiche.code), ("OF", f"{of.code} — {of.label}"), (tr("Statut"), fiche.get_statut_display()),
             (tr("Outillage réservé"), fiche.outillage), (tr("Consignes"), fiche.consignes),
             (tr("Validée par"), (fiche.valide_par.get_full_name() or fiche.valide_par.email) if fiche.valide_par_id else "—")]
    tableau = [[tr("Article"), tr("Désignation"), tr("Quantité"), tr("Disponibilité")]] + [
        [b.article.code, b.article.label[:60], str(b.quantite), b.get_statut_display()]
        for b in of.besoins_matiere.select_related("article")
    ]
    return _pdf(tr("Fiche de préparation %(fiche_code)s") % {"fiche_code": fiche.code}, blocs, tableau if len(tableau) > 1 else None)


def rapport_pdf(rapport):
    of = rapport.ordre
    blocs = [(rapport.get_type_display(), rapport.code), ("OF", f"{of.code} — {of.label}"), (tr("Date"), _d(rapport.date)),
             (tr("Quantité réalisée"), rapport.quantite_realisee), (tr("Quantité rebutée"), rapport.quantite_rebut),
             (tr("Avancement déclaré"), f"{rapport.avancement} %" if rapport.avancement is not None else "—"),
             (tr("Observations"), rapport.observations), (tr("Statut"), rapport.get_statut_display())]
    return _pdf(f"{rapport.get_type_display()} {rapport.code}", blocs)


def transfert_pdf(bt):
    blocs = [(bt.get_type_display(), bt.code), ("OF", f"{bt.ordre.code} — {bt.ordre.label}"),
             (tr("Origine"), bt.origine), (tr("Destination"), bt.destination), (tr("Transporteur"), bt.transporteur),
             (tr("Statut"), bt.get_statut_display()), (tr("Expédié le"), _d(bt.expedie_le)), (tr("Réceptionné le"), _d(bt.receptionne_le)),
             (tr("Réserves"), bt.reserves), (tr("Signatures"), tr("Expéditeur : ____________    Réceptionnaire : ____________"))]
    tableau = [["#", tr("Désignation"), tr("Quantité")]] + [
        [str(i + 1), str(l.get("designation", "")), str(l.get("quantite", ""))] for i, l in enumerate(bt.lignes or [])
    ]
    return _pdf(f"{bt.get_type_display()} {bt.code}", blocs, tableau if len(tableau) > 1 else None)


def rapport_execution_of_pdf(of):
    """Rapport d'exécution consolidé à la terminaison de l'OF (archivé en GED)."""
    rapports = of.rapports.filter(statut=RapportExecution.Statut.VALIDE)
    realise = rapports.aggregate(t=Sum("quantite_realisee"))["t"] or ZERO
    rebut = rapports.aggregate(t=Sum("quantite_rebut"))["t"] or ZERO
    pointe = of.pointages.aggregate(t=Sum("hours"))["t"] or ZERO
    blocs = [("OF", f"{of.code} — {of.label}"), (tr("Affaire"), of.affaire.code if of.affaire_id else "—"),
             (tr("Terminé le"), _d(timezone.localdate())), (tr("Quantité commandée"), of.quantity),
             (tr("Quantité réalisée"), realise), (tr("Quantité rebutée"), rebut),
             (tr("Heures prévues / pointées"), f"{of.calculated_hours} h / {pointe} h")]
    tableau = [[tr("Rapport"), tr("Date"), tr("Réalisé"), tr("Rebut"), tr("Observations")]] + [
        [r.code, _d(r.date), str(r.quantite_realisee), str(r.quantite_rebut), r.observations[:70]] for r in rapports
    ]
    return _pdf(tr("Rapport d'exécution %(of_code)s") % {"of_code": of.code}, blocs, tableau if len(tableau) > 1 else None)


def archiver_rapport_execution(of, user=None):
    from hse.services import archiver_ged

    return archiver_ged(
        en_francais(rapport_execution_of_pdf, of), f"rapport-execution-{of.code}.pdf", f"Rapport d'exécution {of.code}",
        "OPS-REX", "Rapports d'exécution (M3)", of.code, "operations.ordrefabrication", of.pk, user=user,
        mots=("rapport d'exécution", "OF"), module="M3", note="Rapport d'exécution scellé à la terminaison de l'OF",
        retention=int(param("RETENTION_OPERATIONS_ANS")),
    )


def bornes_periode(periode, ref):
    if periode == "mensuel":
        debut = ref.replace(day=1)
        fin = (debut + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    else:
        debut = ref - timedelta(days=ref.weekday())
        fin = debut + timedelta(days=6)
    return debut, fin


def rapport_periodique(periode, ref, voir_montants=True):
    """Rapport hebdomadaire / mensuel consolidé (CODIR) : OF, heures, indicateurs, avancement et marges, alertes."""
    from .cloture import alertes_surcharge
    from .models import BonTransfert, SituationTravaux

    debut, fin = bornes_periode(periode, ref)
    ind = indicateurs(debut, fin)
    suivi = avancement_et_marges(voir_montants)
    pointages = PointageChantier.objects.filter(date__gte=debut, date__lte=fin)
    situations = SituationTravaux.objects.filter(period_end__gte=debut, period_end__lte=fin)
    return {
        "periode": periode, "debut": debut.isoformat(), "fin": fin.isoformat(),
        "synthese": {
            "of_actifs": OrdreFabrication.objects.filter(status__in=ACTIFS).count(),
            "of_a_echeance": OrdreFabrication.objects.filter(planned_end__gte=debut, planned_end__lte=fin).count(),
            "heures_pointees": _s(pointages.aggregate(t=Sum("hours"))["t"] or ZERO),
            "transferts_expedies": BonTransfert.objects.filter(expedie_le__date__gte=debut, expedie_le__date__lte=fin).count(),
            "situations": situations.count(),
            "situations_montant": _s(situations.aggregate(t=Sum("amount"))["t"] or ZERO) if voir_montants else None,
        },
        "indicateurs": ind,
        "affaires": suivi["affaires"],
        "alertes": [a for a in alertes_surcharge() if debut.isoformat() <= a["week"] <= fin.isoformat()],
    }


LIBELLES_SYNTHESE = (
    ("of_actifs", "OF actifs"), ("of_a_echeance", "OF à échéance sur la période"), ("heures_pointees", "Heures pointées"),
    ("transferts_expedies", "BL / transferts expédiés"), ("situations", "Situations de travaux"),
    ("situations_montant", "Montant des situations"),
)
LIBELLES_INDICATEURS = (
    ("quantite_realisee", "Quantité réalisée"), ("quantite_rebut", "Quantité rebutée"), ("taux_rebut_pct", "Taux de rebut (%)"),
    ("productivite_pct", "Productivité (%)"), ("heures_directes", "Heures MO directe"),
    ("heures_indirectes", "Heures MO indirecte"), ("of_en_retard", "OF en retard"), ("ruptures_ouvertes", "Ruptures matière ouvertes"),
)


def _tableaux_rapport(data):
    synthese = [["Rubrique", "Valeur"]] + [[l, "—" if data["synthese"][k] is None else str(data["synthese"][k])] for k, l in LIBELLES_SYNTHESE]
    synthese += [[l, "—" if data["indicateurs"][k] is None else str(data["indicateurs"][k])] for k, l in LIBELLES_INDICATEURS]
    affaires = [["Affaire", "Avancement %", "Coûts engagés", "Coût à terminaison", "Marge à terminaison"]] + [
        [a["affaire_code"], a["avancement_pct"] or "—", a["couts_engages"] or "—", a["cout_a_terminaison"] or "—",
         a["marge_a_terminaison"] or "—"] for a in data["affaires"]
    ]
    alertes = [["Périmètre", "Semaine", "Ressource", "Planifié (h)", "Capacité (h)", "Surcharge %"]] + [
        [a["scope"], a["week"], a.get("resource", "rh"), str(a["planned_hours"]), str(a["capacity"]), str(a["overload_pct"])]
        for a in data["alertes"]
    ]
    return (("Synthèse et indicateurs", synthese), ("Avancement et marge par affaire", affaires), ("Alertes de surcharge", alertes))


def rapport_periodique_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    titre = tr("Rapport %(periode)s des opérations — %(debut)s au %(fin)s") % {"periode": data['periode'], "debut": data['debut'], "fin": data['fin']}
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, title=titre)
    story = [Paragraph(titre, styles["Title"])]
    for sous_titre, tableau in _tableaux_rapport(data):
        story.append(Paragraph(sous_titre, styles["Heading2"]))
        if len(tableau) > 1:
            t = Table(tableau)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12308f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(t)
        else:
            story.append(Paragraph(tr("Néant."), styles["Normal"]))
        story.append(Spacer(1, 4 * mm))
    doc.build(story)
    return buf.getvalue()


def rapport_periodique_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for sous_titre, tableau in _tableaux_rapport(data):
        ws = wb.create_sheet(sous_titre[:31])
        for ligne in tableau:
            ws.append(ligne)
        for cell in ws[1]:
            cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
