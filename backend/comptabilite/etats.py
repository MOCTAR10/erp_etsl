"""États comptables M10 (SYSCOHADA révisé) calculés sur les écritures comptabilisées du noyau.

Grand livre, balance générale, bilan, compte de résultat, balance âgée 401 / 411, plan comptable multi-niveaux,
liasse fiscale et état 2460 (gabarits provisoires), piste d'audit / conservation. Chaque état renvoie la même
structure : `{"titre", "sous_titre", "provisoire", "avertissement", "tableaux": [{"titre", "colonnes", "lignes"}]}`,
exportable en XLSX ou PDF (`?export=xlsx|pdf`). Les rubriques de bilan / résultat suivent les comptes principaux
SYSCOHADA (2 chiffres) : gabarit PROVISOIRE tant qu'ETSL n'a pas fourni son plan détaillé et ses modèles d'états.
"""

import io
from datetime import date, timedelta
from decimal import Decimal

from django.db.models import Q, Sum
from django.utils import timezone

from accounting_kernel.models import AccountMove, AccountMoveLine, FiscalYear
from referentiels.models import Account

ZERO = Decimal("0")
CENT = Decimal("0.01")
COMPTABILISEES = (AccountMove.Status.POSTED, AccountMove.Status.REVERSED)

CLASSES = {
    1: "Comptes de ressources durables", 2: "Comptes d'actif immobilisé", 3: "Comptes de stocks",
    4: "Comptes de tiers", 5: "Comptes de trésorerie", 6: "Comptes de charges des activités ordinaires",
    7: "Comptes de produits des activités ordinaires", 8: "Comptes des autres charges et autres produits",
}

# Comptes principaux SYSCOHADA révisé (niveau 2 du plan).
COMPTES_PRINCIPAUX = {
    "10": "Capital", "11": "Réserves", "12": "Report à nouveau", "13": "Résultat net de l'exercice",
    "14": "Subventions d'investissement", "15": "Provisions réglementées", "16": "Emprunts et dettes assimilées",
    "17": "Dettes de location-acquisition", "18": "Dettes liées à des participations",
    "19": "Provisions pour risques et charges", "20": "Charges immobilisées", "21": "Immobilisations incorporelles",
    "22": "Terrains", "23": "Bâtiments, installations et agencements", "24": "Matériel, mobilier et actifs biologiques",
    "25": "Avances et acomptes versés sur immobilisations", "26": "Titres de participation",
    "27": "Autres immobilisations financières", "28": "Amortissements", "29": "Dépréciations des immobilisations",
    "31": "Marchandises", "32": "Matières premières et fournitures liées", "33": "Autres approvisionnements",
    "34": "Produits en cours", "35": "Services en cours", "36": "Produits finis", "37": "Produits intermédiaires",
    "38": "Stocks en cours de route", "39": "Dépréciations des stocks", "40": "Fournisseurs et comptes rattachés",
    "41": "Clients et comptes rattachés", "42": "Personnel", "43": "Organismes sociaux",
    "44": "État et collectivités publiques", "45": "Organismes internationaux", "46": "Associés et groupe",
    "47": "Débiteurs et créditeurs divers", "48": "Créances et dettes hors activités ordinaires",
    "49": "Dépréciations des comptes de tiers", "50": "Titres de placement", "51": "Valeurs à encaisser",
    "52": "Banques", "53": "Établissements financiers", "54": "Instruments de trésorerie",
    "56": "Banques, crédits de trésorerie", "57": "Caisse", "58": "Régies d'avances et virements internes",
    "59": "Dépréciations des comptes de trésorerie", "60": "Achats et variations de stocks", "61": "Transports",
    "62": "Services extérieurs A", "63": "Services extérieurs B", "64": "Impôts et taxes", "65": "Autres charges",
    "66": "Charges de personnel", "67": "Frais financiers et charges assimilées",
    "68": "Dotations aux amortissements", "69": "Dotations aux provisions", "70": "Ventes",
    "71": "Subventions d'exploitation", "72": "Production immobilisée", "73": "Variations des stocks de biens produits",
    "75": "Autres produits", "77": "Revenus financiers", "78": "Transferts de charges", "79": "Reprises de provisions",
    "81": "Valeurs comptables des cessions d'immobilisations", "82": "Produits des cessions d'immobilisations",
    "83": "Charges hors activités ordinaires", "84": "Produits hors activités ordinaires",
    "85": "Dotations hors activités ordinaires", "86": "Reprises hors activités ordinaires",
    "87": "Participation des travailleurs", "88": "Subventions d'équilibre", "89": "Impôts sur le résultat",
}

AVERTISSEMENT_GABARIT = (
    "Gabarit provisoire (comptes principaux SYSCOHADA) — à confirmer ou modifier par ETSL "
    "(plan comptable détaillé et modèles d'états officiels)."
)


def _col(key, label, montant=False):
    return {"key": key, "label": label, "montant": montant}


def _q(v):
    return (v or ZERO).quantize(CENT)


def lignes_comptabilisees():
    return AccountMoveLine.objects.filter(move__status__in=COMPTABILISEES)


def bornes(date_from=None, date_to=None):
    """Période par défaut : 1er janvier de l'année en cours → aujourd'hui."""
    aujourdhui = timezone.localdate()
    return date_from or date(aujourdhui.year, 1, 1), date_to or aujourdhui


def _soldes_par_compte(qs):
    return {
        r["account_id"]: (r["d"] or ZERO, r["c"] or ZERO)
        for r in qs.values("account_id").annotate(d=Sum("debit"), c=Sum("credit"))
    }


def _gabarit_provisoire(code):
    from .fiscal import get_gabarit

    return get_gabarit(code).est_provisoire


# --------------------------------------------------------------------------------------------------
# Balance générale et grand livre
# --------------------------------------------------------------------------------------------------
def balance(date_from=None, date_to=None, classe=None):
    date_from, date_to = bornes(date_from, date_to)
    base = lignes_comptabilisees()
    if classe:
        base = base.filter(account__account_class=classe)
    ouverture = _soldes_par_compte(base.filter(move__date__lt=date_from))
    mouvements = _soldes_par_compte(base.filter(move__date__gte=date_from, move__date__lte=date_to))
    comptes = Account.objects.filter(pk__in=set(ouverture) | set(mouvements)).order_by("code")
    lignes, tot = [], dict.fromkeys(("ouv_d", "ouv_c", "mvt_d", "mvt_c", "sol_d", "sol_c"), ZERO)
    for compte in comptes:
        od, oc = ouverture.get(compte.pk, (ZERO, ZERO))
        md, mc = mouvements.get(compte.pk, (ZERO, ZERO))
        ouv, fin = od - oc, od - oc + md - mc
        row = {
            "compte": compte.code, "libelle": compte.label, "classe": compte.account_class,
            "ouv_d": _q(max(ouv, ZERO)), "ouv_c": _q(max(-ouv, ZERO)), "mvt_d": _q(md), "mvt_c": _q(mc),
            "sol_d": _q(max(fin, ZERO)), "sol_c": _q(max(-fin, ZERO)),
        }
        for k in tot:
            tot[k] += row[k]
        lignes.append(row)
    lignes.append({"compte": "", "libelle": "TOTAL", "classe": "", **{k: _q(v) for k, v in tot.items()}, "total": True})
    return {
        "code": "balance", "titre": "Balance générale",
        "sous_titre": f"Du {date_from:%d/%m/%Y} au {date_to:%d/%m/%Y}",
        "provisoire": False, "avertissement": "",
        "equilibre": tot["mvt_d"] == tot["mvt_c"] and tot["sol_d"] == tot["sol_c"],
        "tableaux": [{
            "titre": "Balance générale",
            "colonnes": [
                _col("compte", "Compte"), _col("libelle", "Intitulé"),
                _col("ouv_d", "Ouverture débit", True), _col("ouv_c", "Ouverture crédit", True),
                _col("mvt_d", "Mouvements débit", True), _col("mvt_c", "Mouvements crédit", True),
                _col("sol_d", "Solde débiteur", True), _col("sol_c", "Solde créditeur", True),
            ],
            "lignes": lignes,
        }],
    }


def grand_livre(date_from=None, date_to=None, compte=None, limite=5000):
    """Grand livre : par compte, solde à nouveau puis chaque ligne avec solde progressif. `compte` = préfixe."""
    date_from, date_to = bornes(date_from, date_to)
    base = lignes_comptabilisees()
    if compte:
        base = base.filter(account__code__startswith=compte)
    ouverture = _soldes_par_compte(base.filter(move__date__lt=date_from))
    lignes_periode = (
        base.filter(move__date__gte=date_from, move__date__lte=date_to)
        .select_related("account", "move__journal", "partner", "analytic_account")
        .order_by("account__code", "move__date", "move__number", "order")[:limite]
    )
    par_compte = {}
    for l in lignes_periode:
        par_compte.setdefault(l.account, []).append(l)
    for compte_obj in Account.objects.filter(pk__in=ouverture).exclude(pk__in=[c.pk for c in par_compte]):
        par_compte.setdefault(compte_obj, [])
    lignes = []
    tot_d = tot_c = ZERO
    for compte_obj in sorted(par_compte, key=lambda a: a.code):
        od, oc = ouverture.get(compte_obj.pk, (ZERO, ZERO))
        solde = od - oc
        lignes.append({
            "date": "", "journal": "", "piece": "", "compte": compte_obj.code,
            "libelle": f"{compte_obj.label} — solde à nouveau", "tiers": "", "analytique": "",
            "debit": _q(max(solde, ZERO)), "credit": _q(max(-solde, ZERO)), "solde": _q(solde), "entete": True,
        })
        sd = sc = ZERO
        for l in par_compte[compte_obj]:
            solde += l.debit - l.credit
            sd += l.debit
            sc += l.credit
            lignes.append({
                "date": l.move.date.isoformat(), "journal": l.move.journal.code, "piece": l.move.number or "",
                "compte": compte_obj.code, "libelle": l.label or l.move.label, "tiers": l.partner.name if l.partner_id else "",
                "analytique": l.analytic_account.code if l.analytic_account_id else "",
                "debit": _q(l.debit), "credit": _q(l.credit), "solde": _q(solde),
            })
        lignes.append({
            "date": "", "journal": "", "piece": "", "compte": compte_obj.code, "libelle": f"Total {compte_obj.code}",
            "tiers": "", "analytique": "", "debit": _q(sd), "credit": _q(sc), "solde": _q(solde), "total": True,
        })
        tot_d += sd
        tot_c += sc
    lignes.append({
        "date": "", "journal": "", "piece": "", "compte": "", "libelle": "TOTAL GÉNÉRAL", "tiers": "", "analytique": "",
        "debit": _q(tot_d), "credit": _q(tot_c), "solde": "", "total": True,
    })
    return {
        "code": "grand_livre", "titre": "Grand livre",
        "sous_titre": f"Du {date_from:%d/%m/%Y} au {date_to:%d/%m/%Y}" + (f" — comptes {compte}*" if compte else ""),
        "provisoire": False, "avertissement": "",
        "tableaux": [{
            "titre": "Grand livre",
            "colonnes": [
                _col("date", "Date"), _col("journal", "Journal"), _col("piece", "Pièce"), _col("compte", "Compte"),
                _col("libelle", "Libellé"), _col("tiers", "Tiers"), _col("analytique", "Analytique"),
                _col("debit", "Débit", True), _col("credit", "Crédit", True), _col("solde", "Solde", True),
            ],
            "lignes": lignes,
        }],
    }


# --------------------------------------------------------------------------------------------------
# Compte de résultat et bilan
# --------------------------------------------------------------------------------------------------
def _par_principal(qs, sens):
    """{compte principal (2 chiffres): montant} ; sens « credit » = crédit − débit."""
    out = {}
    for r in qs.values("account__code").annotate(d=Sum("debit"), c=Sum("credit")):
        d, c = r["d"] or ZERO, r["c"] or ZERO
        cle = r["account__code"][:2]
        out[cle] = out.get(cle, ZERO) + ((c - d) if sens == "credit" else (d - c))
    return out


def _resultat(qs):
    agg = qs.filter(account__account_class__in=(6, 7, 8)).aggregate(d=Sum("debit"), c=Sum("credit"))
    return (agg["c"] or ZERO) - (agg["d"] or ZERO)


def compte_resultat(date_from=None, date_to=None):
    date_from, date_to = bornes(date_from, date_to)
    qs = lignes_comptabilisees().filter(move__date__gte=date_from, move__date__lte=date_to)
    produits_codes = Q(account__account_class=7) | Q(account__account_class=8, account__account_type=Account.AccountType.INCOME)
    charges_codes = Q(account__account_class=6) | Q(account__account_class=8) & ~Q(account__account_type=Account.AccountType.INCOME)
    produits = _par_principal(qs.filter(produits_codes), "credit")
    charges = _par_principal(qs.filter(charges_codes), "debit")

    def section(titre, donnees):
        rows = [{"rubrique": k, "libelle": COMPTES_PRINCIPAUX.get(k, f"Comptes {k}"), "montant": _q(v)} for k, v in sorted(donnees.items())]
        total = sum(donnees.values(), ZERO)
        rows.append({"rubrique": "", "libelle": f"Total {titre.lower()}", "montant": _q(total), "total": True})
        return rows, total

    lignes_p, total_p = section("Produits", produits)
    lignes_c, total_c = section("Charges", charges)
    resultat = total_p - total_c
    cols = [_col("rubrique", "Rubrique"), _col("libelle", "Libellé"), _col("montant", "Montant", True)]
    return {
        "code": "compte_resultat", "titre": "Compte de résultat",
        "sous_titre": f"Du {date_from:%d/%m/%Y} au {date_to:%d/%m/%Y}",
        "provisoire": _gabarit_provisoire("RESULTAT"), "avertissement": AVERTISSEMENT_GABARIT,
        "resultat": _q(resultat),
        "tableaux": [
            {"titre": "Produits", "colonnes": cols, "lignes": lignes_p},
            {"titre": "Charges", "colonnes": cols, "lignes": lignes_c},
            {"titre": "Résultat", "colonnes": cols, "lignes": [{
                "rubrique": "13", "libelle": "Résultat net (bénéfice si positif, perte si négatif)",
                "montant": _q(resultat), "total": True,
            }]},
        ],
    }


BILAN_ACTIF = (
    ("AD", "Immobilisations incorporelles (nettes)", ("20", "21")),
    ("AI", "Immobilisations corporelles (nettes)", ("22", "23", "24", "25", "28", "29")),
    ("AQ", "Immobilisations financières", ("26", "27")),
    ("BB", "Stocks et en-cours", ("31", "32", "33", "34", "35", "36", "37", "38", "39")),
)
BILAN_PASSIF = (
    ("CA", "Capital", ("10",)),
    ("CD", "Réserves et report à nouveau", ("11", "12")),
    ("CL", "Subventions et provisions réglementées", ("14", "15")),
    ("DA", "Emprunts et dettes financières", ("16", "17", "18")),
    ("DD", "Provisions pour risques et charges", ("19",)),
)


def bilan(date_to=None):
    """Bilan à une date (cumul de toutes les écritures comptabilisées). Comptes de tiers et de trésorerie
    classés à l'actif ou au passif selon le sens de leur solde ; résultat = classes 6, 7, 8 non affectées."""
    date_to = date_to or timezone.localdate()
    qs = lignes_comptabilisees().filter(move__date__lte=date_to)
    soldes = {}
    for r in qs.values("account__code").annotate(d=Sum("debit"), c=Sum("credit")):
        soldes[r["account__code"]] = (r["d"] or ZERO) - (r["c"] or ZERO)
    actif, passif = [], []

    def somme(prefixes, signe):
        return sum((s for code, s in soldes.items() if code[:2] in prefixes), ZERO) * signe

    for code, libelle, prefixes in BILAN_ACTIF:
        actif.append({"rubrique": code, "libelle": libelle, "montant": _q(somme(prefixes, 1))})
    tiers_d = sum((s for code, s in soldes.items() if code[0] == "4" and s > 0), ZERO)
    tiers_c = sum((-s for code, s in soldes.items() if code[0] == "4" and s < 0), ZERO)
    treso_d = sum((s for code, s in soldes.items() if code[0] == "5" and s > 0), ZERO)
    treso_c = sum((-s for code, s in soldes.items() if code[0] == "5" and s < 0), ZERO)
    actif += [
        {"rubrique": "BH", "libelle": "Créances et emplois assimilés (comptes de tiers débiteurs)", "montant": _q(tiers_d)},
        {"rubrique": "BT", "libelle": "Trésorerie-Actif", "montant": _q(treso_d)},
    ]
    for code, libelle, prefixes in BILAN_PASSIF:
        passif.append({"rubrique": code, "libelle": libelle, "montant": _q(somme(prefixes, -1))})
    resultat = _resultat(qs)
    passif += [
        {"rubrique": "CJ", "libelle": "Résultat net de l'exercice (non affecté)", "montant": _q(resultat)},
        {"rubrique": "DP", "libelle": "Passif circulant (comptes de tiers créditeurs)", "montant": _q(tiers_c)},
        {"rubrique": "DT", "libelle": "Trésorerie-Passif", "montant": _q(treso_c)},
    ]
    # Compte 13 (résultat porté au bilan par une écriture) : rubrique distincte s'il est mouvementé.
    if somme(("13",), -1):
        passif.insert(len(BILAN_PASSIF), {"rubrique": "CH", "libelle": "Résultat affecté (compte 13)", "montant": _q(somme(("13",), -1))})
    total_actif = sum((r["montant"] for r in actif), ZERO)
    total_passif = sum((r["montant"] for r in passif), ZERO)
    actif.append({"rubrique": "BZ", "libelle": "TOTAL ACTIF", "montant": _q(total_actif), "total": True})
    passif.append({"rubrique": "DZ", "libelle": "TOTAL PASSIF", "montant": _q(total_passif), "total": True})
    cols = [_col("rubrique", "Réf."), _col("libelle", "Rubrique"), _col("montant", "Montant net", True)]
    return {
        "code": "bilan", "titre": "Bilan", "sous_titre": f"Au {date_to:%d/%m/%Y}",
        "provisoire": _gabarit_provisoire("BILAN"), "avertissement": AVERTISSEMENT_GABARIT,
        "equilibre": total_actif == total_passif,
        "tableaux": [
            {"titre": "Actif", "colonnes": cols, "lignes": actif},
            {"titre": "Passif", "colonnes": cols, "lignes": passif},
        ],
    }


# --------------------------------------------------------------------------------------------------
# Balance âgée clients (411) / fournisseurs (401)
# --------------------------------------------------------------------------------------------------
def balance_agee(nature="clients", date_to=None):
    from .services import get_param

    date_to = date_to or timezone.localdate()
    prefixe, signe = ("411", 1) if nature == "clients" else ("401", -1)
    t1, t2, t3 = (int(get_param(c)) for c in ("BAL_AGEE_TRANCHE_1", "BAL_AGEE_TRANCHE_2", "BAL_AGEE_TRANCHE_3"))
    tranches = [
        ("t0", f"0 – {t1} j", 0, t1), ("t1", f"{t1 + 1} – {t2} j", t1 + 1, t2),
        ("t2", f"{t2 + 1} – {t3} j", t2 + 1, t3), ("t3", f"> {t3} j", t3 + 1, None),
    ]
    qs = (
        lignes_comptabilisees().filter(account__code__startswith=prefixe, move__date__lte=date_to, lettrage_actif__isnull=True)
        .select_related("partner", "move")
    )
    par_tiers = {}
    for l in qs:
        cle = l.partner_id
        row = par_tiers.setdefault(cle, {
            "tiers": l.partner.name if l.partner_id else "(sans tiers)", "code": l.partner.code if l.partner_id else "",
            **{k: ZERO for k, *_ in tranches}, "total": ZERO,
        })
        montant = (l.debit - l.credit) * signe
        age = (date_to - l.move.date).days
        for k, _label, mini, maxi in tranches:
            if age >= mini and (maxi is None or age <= maxi):
                row[k] += montant
                break
        row["total"] += montant
    lignes = [
        {**r, **{k: _q(r[k]) for k, *_ in tranches}, "total": _q(r["total"])}
        for r in sorted(par_tiers.values(), key=lambda r: r["tiers"]) if r["total"]
    ]
    totaux = {k: _q(sum((r[k] for r in lignes), ZERO)) for k, *_ in tranches}
    lignes.append({"tiers": "TOTAL", "code": "", **totaux, "total": _q(sum(totaux.values(), ZERO)), "total_ligne": True})
    from .services import parametres_provisoires

    return {
        "code": "balance_agee", "titre": f"Balance âgée {'clients (411)' if nature == 'clients' else 'fournisseurs (401)'}",
        "sous_titre": f"Au {date_to:%d/%m/%Y} — lignes non lettrées, ancienneté depuis la date de pièce",
        "provisoire": parametres_provisoires(["BAL_AGEE_TRANCHE_1", "BAL_AGEE_TRANCHE_2", "BAL_AGEE_TRANCHE_3"]),
        "avertissement": "Tranches d'ancienneté provisoires — à confirmer ou modifier par ETSL (Paramètres).",
        "tableaux": [{
            "titre": "Balance âgée",
            "colonnes": [_col("code", "Code"), _col("tiers", "Tiers")]
            + [_col(k, label, True) for k, label, *_ in tranches] + [_col("total", "Solde", True)],
            "lignes": lignes,
        }],
    }


# --------------------------------------------------------------------------------------------------
# Plan comptable multi-niveaux et piste d'audit / conservation
# --------------------------------------------------------------------------------------------------
def plan_comptable():
    """Plan à 4 niveaux : classe → compte principal (2 chiffres) → compte → sous-compte (parent renseigné)."""
    comptes = list(Account.objects.select_related("parent").order_by("code"))
    racines = {a.pk: a.code.rstrip("0") for a in comptes}

    def parent_de(a):
        """Parent renseigné, sinon le compte le plus proche dont `a` prolonge le numéro (401100 → 401000)."""
        if a.parent_id:
            return a.parent_id
        meilleur = None
        for b in comptes:
            r = racines[b.pk]
            if b.pk != a.pk and r and len(r) < len(racines[a.pk]) and racines[a.pk].startswith(r):
                if meilleur is None or len(r) > len(racines[meilleur]):
                    meilleur = b.pk
        return meilleur

    parents = {a.pk: parent_de(a) for a in comptes}
    lignes, vus_classe, vus_principal = [], set(), set()
    for a in comptes:
        if parents[a.pk]:
            continue
        if a.account_class not in vus_classe:
            vus_classe.add(a.account_class)
            lignes.append({"niveau": 1, "code": str(a.account_class), "libelle": CLASSES.get(a.account_class, ""), "parent": ""})
        principal = a.code[:2]
        if principal not in vus_principal:
            vus_principal.add(principal)
            lignes.append({"niveau": 2, "code": principal, "libelle": COMPTES_PRINCIPAUX.get(principal, ""), "parent": str(a.account_class)})
        lignes.append({"niveau": 3, "code": a.code, "libelle": a.label, "parent": principal})
        for enfant in sorted((c for c in comptes if parents[c.pk] == a.pk), key=lambda c: c.code):
            lignes.append({"niveau": 4, "code": enfant.code, "libelle": enfant.label, "parent": a.code})
    return {
        "code": "plan_comptable", "titre": "Plan comptable SYSCOHADA (multi-niveaux)",
        "sous_titre": f"{len(comptes)} comptes — classe › compte principal › compte › sous-compte",
        "provisoire": True,
        "avertissement": "Plan provisoire : comptes détaillés et nombre de niveaux à confirmer par ETSL.",
        "tableaux": [{
            "titre": "Plan comptable",
            "colonnes": [_col("niveau", "Niveau"), _col("code", "Compte"), _col("libelle", "Intitulé"), _col("parent", "Rattaché à")],
            "lignes": lignes,
        }],
    }


def conservation():
    """Piste d'audit : écritures conservées par exercice (immuables, extournes tracées) et date de fin de conservation."""
    from documents.models import Document

    from .fiscal import TYPES_FISCAUX
    from .services import get_param

    ans = int(get_param("CONSERVATION_ANS"))
    lignes = []
    for fy in FiscalYear.objects.order_by("-year"):
        moves = AccountMove.objects.filter(period__fiscal_year=fy)
        lignes.append({
            "exercice": fy.year, "statut": fy.get_status_display(),
            "ecritures": moves.filter(status__in=COMPTABILISEES).count(),
            "extournees": moves.filter(status=AccountMove.Status.REVERSED).count(),
            "manuelles": moves.filter(status__in=COMPTABILISEES, source=AccountMove.SOURCE_MANUELLE).count(),
            "pieces_ged": Document.objects.filter(type__code__in=TYPES_FISCAUX, document_date__year=fy.year).count(),
            "conserver_jusqu_au": (date(fy.end_date.year + ans, fy.end_date.month, fy.end_date.day)).isoformat(),
        })
    from .services import parametres_provisoires

    return {
        "code": "conservation", "titre": "Piste d'audit et conservation des écritures",
        "sous_titre": f"Écritures comptabilisées immuables (correction par extourne), conservées {ans} ans après la clôture",
        "provisoire": parametres_provisoires(["CONSERVATION_ANS"]),
        "avertissement": "Durée de conservation provisoire — à confirmer par ETSL.",
        "tableaux": [{
            "titre": "Conservation par exercice",
            "colonnes": [
                _col("exercice", "Exercice"), _col("statut", "Statut"), _col("ecritures", "Écritures"),
                _col("extournees", "Extournées"), _col("manuelles", "Saisies manuelles"),
                _col("pieces_ged", "Pièces fiscales en GED"), _col("conserver_jusqu_au", "Conserver jusqu'au"),
            ],
            "lignes": lignes,
        }],
    }


# --------------------------------------------------------------------------------------------------
# Exports XLSX / PDF communs à tous les états
# --------------------------------------------------------------------------------------------------
def jsonable(etat):
    """Décimaux en chaînes (Response brute) ; masque `None` conservé."""
    def conv(v):
        if isinstance(v, Decimal):
            return str(v)
        if isinstance(v, (date,)):
            return v.isoformat()
        if isinstance(v, dict):
            return {k: conv(x) for k, x in v.items()}
        if isinstance(v, list):
            return [conv(x) for x in v]
        return v

    return conv(etat)


def _fmt(v):
    if isinstance(v, Decimal):
        return f"{v:,.0f}".replace(",", " ")
    return "" if v is None else str(v)


def to_xlsx(etat):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for i, tab in enumerate(etat["tableaux"]):
        ws = wb.create_sheet(title=(tab["titre"] or f"Feuille {i + 1}")[:31].replace("/", "-"))
        ws.append([etat["titre"]])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([etat.get("sous_titre", "")])
        if etat.get("provisoire"):
            ws.append([f"PROVISOIRE — {etat.get('avertissement', '')}"])
        ws.append([])
        ws.append([c["label"] for c in tab["colonnes"]])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for row in tab["lignes"]:
            values = []
            for c in tab["colonnes"]:
                v = row.get(c["key"])
                values.append(float(v) if isinstance(v, Decimal) else ("" if v is None else v))
            ws.append(values)
            if row.get("total") or row.get("total_ligne") or row.get("entete"):
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
        for col_cells in ws.iter_cols(min_row=5):
            width = max((len(str(c.value or "")) for c in col_cells), default=8) + 2
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width, 10), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(etat):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=etat["titre"])
    story = [Paragraph(f"ETSL — {etat['titre']}", styles["Title"]), Paragraph(etat.get("sous_titre", ""), styles["Normal"])]
    if etat.get("provisoire"):
        story.append(Paragraph(f"<font color='#d4142a'><b>PROVISOIRE</b> — {etat.get('avertissement', '')}</font>", styles["Normal"]))
    for tab in etat["tableaux"]:
        story += [Spacer(1, 4 * mm), Paragraph(f"<b>{tab['titre']}</b>", styles["Heading3"])]
        data = [[c["label"] for c in tab["colonnes"]]]
        gras = []
        for i, row in enumerate(tab["lignes"], start=1):
            data.append([Paragraph(_fmt(row.get(c["key"])), styles["BodyText"]) if not c["montant"] else _fmt(row.get(c["key"]))
                         for c in tab["colonnes"]])
            if row.get("total") or row.get("total_ligne") or row.get("entete"):
                gras.append(i)
        table = Table(data, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12308f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
        for j, c in enumerate(tab["colonnes"]):
            if c["montant"]:
                style.append(("ALIGN", (j, 1), (j, -1), "RIGHT"))
        for i in gras:
            style += [("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"), ("BACKGROUND", (0, i), (-1, i), colors.whitesmoke)]
        table.setStyle(TableStyle(style))
        story.append(table)
    story += [Spacer(1, 4 * mm), Paragraph(f"Édité le {timezone.localtime():%d/%m/%Y %H:%M} — ERP ETSL", styles["Italic"])]
    doc.build(story)
    return buf.getvalue()


def semaine(jour):
    """Lundi de la semaine de `jour`."""
    return jour - timedelta(days=jour.weekday())
